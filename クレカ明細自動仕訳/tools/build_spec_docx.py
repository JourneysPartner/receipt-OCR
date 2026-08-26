from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent.parent
SOURCE = ROOT / "credit_card_import_normalization_system_spec_v2.0.md"
OUTPUT = ROOT / "credit_card_import_normalization_system_spec_v2.0.docx"
SKILL_SCRIPTS = Path(
    r"C:\Users\ko-ch\.codex\plugins\cache\openai-primary-runtime\documents\26.819.11345\skills\documents\scripts"
)
sys.path.insert(0, str(SKILL_SCRIPTS))
from table_geometry import apply_table_geometry, audit_docx_tables  # noqa: E402


# compact_reference_guide with a named Japanese-font override.
FONT = "Yu Gothic"
MONO = "Consolas"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "1D2733"
MUTED = "667085"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
CODE_FILL = "F5F7FA"
BORDER = "C8D0DA"
TABLE_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120
CELL_MARGINS = {"top": 90, "bottom": 90, "start": 120, "end": 120}


def set_run_font(run, name=FONT, size=None, color=None, bold=None, italic=None):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        rfonts.set(qn(f"w:{attr}"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_borders(cell, color=BORDER, size="4"):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_borders = tc_pr.find(qn("w:tcBorders"))
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        el = tc_borders.find(qn(tag))
        if el is None:
            el = OxmlElement(tag)
            tc_borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), size)
        el.set(qn("w:color"), color)


def set_paragraph_shading(paragraph, fill):
    ppr = paragraph._p.get_or_add_pPr()
    shd = ppr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        ppr.append(shd)
    shd.set(qn("w:fill"), fill)


def add_page_field(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])
    set_run_font(run, size=9, color=MUTED)


def configure_styles(doc: Document):
    styles = doc.styles

    normal = styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.2

    for style_name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 18, 9),
        ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 11.5, DARK_BLUE, 10, 5),
    ):
        style = styles[style_name]
        style.font.name = FONT
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True

    for name in ("List Bullet", "List Number"):
        style = styles[name]
        style.font.name = FONT
        style.font.size = Pt(10.5)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        style.paragraph_format.left_indent = Inches(0.375)
        style.paragraph_format.first_line_indent = Inches(-0.188)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.2

    if "Code Block" not in styles:
        code_style = styles.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
    else:
        code_style = styles["Code Block"]
    code_style.font.name = MONO
    code_style.font.size = Pt(8.5)
    code_style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    code_style.paragraph_format.left_indent = Inches(0.22)
    code_style.paragraph_format.right_indent = Inches(0.16)
    code_style.paragraph_format.space_before = Pt(4)
    code_style.paragraph_format.space_after = Pt(7)
    code_style.paragraph_format.line_spacing = 1.0

    if "TOC Entry" not in styles:
        toc = styles.add_style("TOC Entry", WD_STYLE_TYPE.PARAGRAPH)
    else:
        toc = styles["TOC Entry"]
    toc.font.name = FONT
    toc.font.size = Pt(9.5)
    toc.font.color.rgb = RGBColor.from_string(INK)
    toc._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    toc.paragraph_format.space_after = Pt(3)
    toc.paragraph_format.left_indent = Inches(0.08)


def configure_page(doc: Document):
    doc.settings.odd_and_even_pages_header_footer = False
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    hp.paragraph_format.space_after = Pt(0)
    # A white spacer preserves a stable header part while keeping the page clean.
    hr = hp.add_run(" ")
    set_run_font(hr, size=8.5, color="FFFFFF")

    footer = section.footer
    fp = footer.paragraphs[0]
    add_page_field(fp)


def add_title_page(doc: Document):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(24)
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run("SYSTEM SPECIFICATION")
    set_run_font(r, size=10, color=BLUE, bold=True)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run("クレジットカード明細\n自動取込・取引先正規化システム")
    set_run_font(r, size=24, color=INK, bold=True)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(24)
    r = p.add_run("仕様書 Ver.2.0")
    set_run_font(r, size=14, color=DARK_BLUE, bold=True)

    metadata = [
        ("作成日", "2026年8月23日"),
        ("対象", "Google Drive／Googleスプレッドシート／Google Apps Script"),
        ("文書状態", "5ラウンドレビュー完了・実装前決定事項あり"),
        ("対象範囲", "顧客提出明細からfreeeクレカ出納帳「入力用シート」まで"),
    ]
    for label, value in metadata:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(3)
        lr = p.add_run(f"{label}：")
        set_run_font(lr, size=10.5, color=INK, bold=True)
        vr = p.add_run(value)
        set_run_font(vr, size=10.5, color=INK)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(24)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.left_indent = Inches(0.16)
    p.paragraph_format.right_indent = Inches(0.16)
    set_paragraph_shading(p, LIGHT_BLUE)
    r = p.add_run(
        "本版は、CodexとClaude Codeによる5ラウンドのレビューを反映し、状態管理、冪等性、"
        "取引先判定、承認、監査、freee取込後の変更管理、および受入条件を統合した実装基準である。"
        "第27章の業務判断を確定してから本番実装へ進む。"
    )
    set_run_font(r, size=10.5, color=DARK_BLUE, bold=True)

    doc.add_page_break()


def parse_inline(paragraph, text: str, *, size=None, color=None, bold=False):
    text = text.replace("  ", " ").strip()
    parts = re.split(r"(`[^`]+`|\*\*[^*]+\*\*)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            set_run_font(run, name=MONO, size=size or 9.5, color=color or DARK_BLUE, bold=bold)
        elif part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            set_run_font(run, size=size, color=color, bold=True)
        else:
            run = paragraph.add_run(part)
            set_run_font(run, size=size, color=color, bold=bold)


def add_toc(doc: Document, lines: list[str]):
    h = doc.add_paragraph(style="Heading 1")
    h.paragraph_format.space_before = Pt(0)
    h.paragraph_format.space_after = Pt(12)
    parse_inline(h, "目次")
    for line in lines:
        if re.match(r"^# \d+\.", line):
            p = doc.add_paragraph(style="TOC Entry")
            parse_inline(p, line[2:])
    doc.add_page_break()


def add_table(doc: Document, rows: list[list[str]]):
    if not rows:
        return
    cols = max(len(row) for row in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    if cols == 2:
        first_max = max(len(row[0]) if row else 0 for row in rows)
        widths = [3600, 5760] if first_max > 18 else [2700, 6660]
    elif cols == 3:
        widths = [1700, 3830, 3830]
    else:
        base = TABLE_WIDTH_DXA // cols
        widths = [base] * cols
        widths[-1] += TABLE_WIDTH_DXA - sum(widths)

    for i, row in enumerate(rows):
        for j in range(cols):
            cell = table.cell(i, j)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_borders(cell)
            if i == 0:
                set_cell_shading(cell, LIGHT_BLUE)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            text = row[j] if j < len(row) else ""
            parse_inline(p, text, size=9, bold=(i == 0), color=(DARK_BLUE if i == 0 else INK))
    set_repeat_header(table.rows[0])
    apply_table_geometry(
        table,
        widths,
        table_width_dxa=TABLE_WIDTH_DXA,
        indent_dxa=TABLE_INDENT_DXA,
        cell_margins_dxa=CELL_MARGINS,
    )
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_code_block(doc: Document, code_lines: list[str]):
    p = doc.add_paragraph(style="Code Block")
    set_paragraph_shading(p, CODE_FILL)
    for idx, line in enumerate(code_lines):
        if idx:
            p.add_run().add_break()
        r = p.add_run(line)
        set_run_font(r, name=MONO, size=8.5, color=INK)


def clean_table_cell(value: str) -> str:
    return value.strip().replace("\\|", "|")


def create_restarted_numbering(doc: Document) -> int:
    numbering = doc.part.numbering_part.element
    style = doc.styles["List Number"]
    style_num_id = int(style._element.pPr.numPr.numId.val)
    abstract_num_id = None
    existing_ids = []
    for num in numbering.findall(qn("w:num")):
        num_id = int(num.get(qn("w:numId")))
        existing_ids.append(num_id)
        if num_id == style_num_id:
            abstract = num.find(qn("w:abstractNumId"))
            abstract_num_id = int(abstract.get(qn("w:val")))
    if abstract_num_id is None:
        raise RuntimeError("List Number abstract numbering definition not found")

    new_num_id = max(existing_ids, default=0) + 1
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(new_num_id))
    abstract = OxmlElement("w:abstractNumId")
    abstract.set(qn("w:val"), str(abstract_num_id))
    num.append(abstract)
    override = OxmlElement("w:lvlOverride")
    override.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:startOverride")
    start.set(qn("w:val"), "1")
    override.append(start)
    num.append(override)
    numbering.append(num)
    return new_num_id


def apply_numbering(paragraph, num_id: int):
    ppr = paragraph._p.get_or_add_pPr()
    num_pr = ppr.get_or_add_numPr()
    ilvl = num_pr.find(qn("w:ilvl"))
    if ilvl is None:
        ilvl = OxmlElement("w:ilvl")
        num_pr.append(ilvl)
    ilvl.set(qn("w:val"), "0")
    num_id_el = num_pr.find(qn("w:numId"))
    if num_id_el is None:
        num_id_el = OxmlElement("w:numId")
        num_pr.append(num_id_el)
    num_id_el.set(qn("w:val"), str(num_id))


def build_body(doc: Document, lines: list[str]):
    start = next(i for i, line in enumerate(lines) if re.match(r"^# 1\.", line))
    i = start
    first_h1 = True
    active_num_id = None
    while i < len(lines):
        raw = lines[i].rstrip()
        stripped = raw.strip()
        if not stripped or stripped == "---":
            i += 1
            continue

        if stripped.startswith("```"):
            active_num_id = None
            code = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code.append(lines[i].rstrip())
                i += 1
            add_code_block(doc, code)
            i += 1
            continue

        if stripped.startswith("|"):
            active_num_id = None
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                table_lines.append(lines[i].strip())
                i += 1
            rows = []
            for idx, tline in enumerate(table_lines):
                cells = [clean_table_cell(v) for v in tline.strip("|").split("|")]
                if idx == 1 and all(re.fullmatch(r":?-{3,}:?", c) for c in cells):
                    continue
                rows.append(cells)
            add_table(doc, rows)
            continue

        heading = re.match(r"^(#{1,3})\s+(.+)$", stripped)
        if heading:
            active_num_id = None
            level = len(heading.group(1))
            text = heading.group(2)
            p = doc.add_paragraph(style=f"Heading {level}")
            if level == 1:
                first_h1 = False
            parse_inline(p, text)
            i += 1
            continue

        checkbox = re.match(r"^-\s+\[([ xX])\]\s+(.+)$", stripped)
        if checkbox:
            active_num_id = None
            p = doc.add_paragraph(style="List Bullet")
            mark = "☒ " if checkbox.group(1).lower() == "x" else "☐ "
            parse_inline(p, mark + checkbox.group(2))
            i += 1
            continue

        bullet = re.match(r"^-\s+(.+)$", stripped)
        if bullet:
            active_num_id = None
            p = doc.add_paragraph(style="List Bullet")
            parse_inline(p, bullet.group(1))
            i += 1
            continue

        numbered = re.match(r"^\d+\.\s+(.+)$", stripped)
        if numbered:
            if active_num_id is None:
                active_num_id = create_restarted_numbering(doc)
            p = doc.add_paragraph(style="List Number")
            apply_numbering(p, active_num_id)
            parse_inline(p, numbered.group(1))
            i += 1
            continue

        active_num_id = None
        p = doc.add_paragraph()
        parse_inline(p, stripped)
        i += 1


def add_core_properties(doc: Document):
    props = doc.core_properties
    props.title = "クレジットカード明細 自動取込・取引先正規化システム 仕様書 Ver.2.0"
    props.subject = "Google Drive提出明細の自動取込・freee取引先正規化"
    props.author = ""
    props.keywords = "クレジットカード, Google Apps Script, freee, 取引先正規化"
    props.comments = "要件確定版"


def main():
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    doc = Document()
    configure_page(doc)
    configure_styles(doc)
    add_core_properties(doc)
    add_title_page(doc)
    add_toc(doc, lines)
    build_body(doc, lines)
    doc.save(OUTPUT)
    failures = audit_docx_tables(OUTPUT)
    if failures:
        raise SystemExit(f"table geometry audit failed: {failures}")
    print(f"created: {OUTPUT}")


if __name__ == "__main__":
    main()
