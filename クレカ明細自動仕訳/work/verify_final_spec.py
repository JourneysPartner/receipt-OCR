import re
from pathlib import Path
from docx import Document

root = Path(r"C:\Users\ko-ch\Desktop\仕事\receipt OCR\クレカ明細自動仕訳")
md_path = root / "credit_card_import_normalization_system_spec_v2.0.md"
docx_path = root / "credit_card_import_normalization_system_spec_v2.0.docx"

md = md_path.read_text(encoding="utf-8")
doc = Document(docx_path)
doc_text = "\n".join(p.text for p in doc.paragraphs)
doc_text += "\n" + "\n".join(
    cell.text for table in doc.tables for row in table.rows for cell in row.cells
)

top_sections = [int(n) for n in re.findall(r"^# (\d+)\. ", md, flags=re.MULTILINE)]
expected_sections = list(range(1, 30))

required = [
    "K列には顧客提出ファイルに記載された元利用店名",
    "freee取引先",
    "partnerResolutionStatus",
    "CUSTOMER_FIX_REQUIRED",
    "CONTINUATION_TRIGGER_CREATE_FAILED",
    "取込時点の各取引のB・F・I・K・M列ハッシュ",
    "オーナー管理者の申請は即時承認",
    "AI結果は必ず人間が確認",
    "未freee取込の転記行、要確認、辞書候補を取り消す最終実行権限が未確定",
    "5ラウンドレビュー完了・実装前決定事項あり",
]

missing_md = [term for term in required if term not in md]
missing_docx = [term for term in required if term not in doc_text]
heading_counts = {
    style: sum(1 for p in doc.paragraphs if p.style.name == style)
    for style in ("Heading 1", "Heading 2", "Heading 3")
}

print(f"markdown_top_sections={top_sections}")
print(f"docx_paragraphs={len(doc.paragraphs)} tables={len(doc.tables)} headings={heading_counts}")
print(f"missing_md={missing_md}")
print(f"missing_docx={missing_docx}")

if top_sections != expected_sections:
    raise SystemExit("Top-level section numbering is not 1..29")
if len(doc.tables) != 17:
    raise SystemExit("Unexpected table count")
if heading_counts["Heading 1"] != 30:
    raise SystemExit("Unexpected Heading 1 count")
if missing_md or missing_docx:
    raise SystemExit("Required content missing")
if re.search(r"\b(?:TODO|TBD)\b|\{\{.*?\}\}", md):
    raise SystemExit("Placeholder token found")
print("final content verification: PASS")
